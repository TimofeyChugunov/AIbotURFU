import base64
import io
import json
import os
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import server


def pdf_bytes():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (Derivative describes a rate of change.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO(); writer.write(output)
    return output.getvalue()


class AppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.data_patch = patch.object(server, 'DATA', Path(cls.temp.name)); cls.data_patch.start()
        cls.env_patch = patch.dict(os.environ, {'AI_MODE': 'demo', 'ADMIN_PASSWORD': 'test-password', 'OPENAI_API_KEY': '', 'OPENAI_MODEL': ''}); cls.env_patch.start()
        server.init_db()
        cls.httpd = server.ThreadingHTTPServer(('127.0.0.1', 0), server.Handler)
        cls.url = f'http://127.0.0.1:{cls.httpd.server_port}'
        cls.thread = threading.Thread(target=cls.httpd.serve_forever, daemon=True); cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown(); cls.httpd.server_close(); cls.thread.join()
        cls.env_patch.stop(); cls.data_patch.stop(); cls.temp.cleanup()

    def setUp(self):
        with server.connect() as db:
            db.execute('DELETE FROM materials')
        server.Handler.sessions.clear(); server.Handler.attempts.clear()
        self.cookie = ''

    def request(self, path, data=None, method=None, headers=None):
        raw = json.dumps(data).encode() if data is not None else None
        req = urllib.request.Request(self.url + path, data=raw, headers={'Content-Type': 'application/json', 'Cookie': self.cookie, **(headers or {})}, method=method)
        try:
            response = urllib.request.urlopen(req, timeout=5)
        except urllib.error.HTTPError as exc:
            response = exc
        with response:
            content = response.read()
            content_type = response.headers.get_content_type()
            result = (json.loads(content) if content_type == 'application/json' else
                      content if response.headers.get('Content-Disposition', '').startswith('attachment') or content_type == 'application/pdf' else content.decode('utf-8'))
            return response.status, result, response.headers

    def login(self):
        status, _, headers = self.request('/api/login', {'password': 'test-password'})
        self.assertEqual(status, 200)
        self.cookie = headers['Set-Cookie'].split(';')[0]

    def upload(self, name='physics.txt', content='Закон Ома: ток равен напряжению, делённому на сопротивление.'.encode()):
        return self.request('/api/materials', {'name': name, 'content': base64.b64encode(content).decode()})

    def test_private_docx_demo_report_is_not_a_shared_material(self):
        from test_docx_review import document_bytes
        attachment = {'name': 'private.docx', 'content': base64.b64encode(document_bytes()).decode()}
        status, result, _ = self.request('/api/chat', {'message': 'Проверь оформление', 'mode': 'report', 'attachment': attachment})
        self.assertEqual(status, 200)
        self.assertTrue(result['demo'])
        self.assertEqual(result['document_review']['name'], 'private.docx')
        self.assertIn('не оценивалось ИИ', result['text'])
        self.assertEqual(self.request('/api/materials')[1], [])
        self.assertEqual(self.request('/assets/private.docx')[0], 404)

    def test_docx_is_passed_as_untrusted_context_to_real_adapter(self):
        from test_docx_review import document_bytes
        attachment = {'name': 'private.docx', 'content': base64.b64encode(document_bytes('PRIVATE-MARKER-729')).decode()}
        with patch.dict(os.environ, {'AI_MODE': 'ollama'}), patch('local_ai.generate', return_value=('Замечание по работе.', [])) as model:
            status, result, _ = self.request('/api/chat', {'message': 'Проверь введение', 'mode': 'report', 'attachment': attachment})
        self.assertEqual(status, 200)
        self.assertFalse(result['demo'])
        self.assertIn('PRIVATE-MARKER-729', model.call_args.args[4])
        self.assertIn('не методичка', model.call_args.args[4])
        self.assertIn('не объявляй параметры нарушением ГОСТ', model.call_args.args[0])
        self.assertIn('document_review', result)
        self.assertEqual(self.request('/api/materials')[1], [])

    def test_bad_chat_attachments_return_clear_errors(self):
        for attachment in [None, {}, {'name': 'x.docx', 'content': '%%%bad'},
                           {'name': 'x.docx', 'content': base64.b64encode(b'not docx').decode()},
                           {'name': 'x.doc', 'content': base64.b64encode(b'old doc').decode()}]:
            status, result, _ = self.request('/api/chat', {'message': 'Проверь', 'mode': 'report', 'attachment': attachment})
            self.assertEqual(status, 400, attachment)
            self.assertIn('error', result)

    def test_static_and_private_files(self):
        for path in ('/', '/admin', '/app.js', '/admin.js', '/styles.css', '/theme.js', '/theme.css'):
            status, content, headers = self.request(path)
            self.assertEqual(status, 200); self.assertTrue(content)
            self.assertIn("default-src 'self'", headers['Content-Security-Policy'])
        for path in ('/server.py', '/.env', '/data/materials.sqlite3', '/../README.md'):
            self.assertEqual(self.request(path)[0], 404)

    def test_demo_never_calls_provider(self):
        with patch('server.urllib.request.urlopen', side_effect=AssertionError('Network forbidden')):
            for mode in server.MODES:
                result = server.answer('Производная', mode, [])
                self.assertTrue(result['demo']); self.assertIn('не ответ ИИ', result['text'])

    def test_admin_required_and_logout(self):
        self.assertEqual(self.upload()[0], 401)
        self.assertEqual(self.request('/api/login', {'password': 'wrong'})[0], 401)
        self.login(); self.assertTrue(self.request('/api/status')[1]['admin'])
        self.assertEqual(self.request('/api/logout', {})[0], 200)
        self.assertEqual(self.upload()[0], 401)

    def test_txt_upload_search_delete(self):
        self.login(); status, file, _ = self.upload(); self.assertEqual(status, 201)
        self.assertEqual(len(self.request('/api/materials')[1]), 1)
        status, result, _ = self.request('/api/chat', {'message': 'Объясни закон Ома', 'mode': 'explain'})
        self.assertEqual(status, 200); self.assertEqual(result['sources'][0]['name'], 'physics.txt')
        self.assertIn('ток равен', result['text'])
        self.assertEqual(self.request('/api/materials/' + file['id'], method='DELETE')[0], 200)
        self.assertEqual(server.search_materials('Закон Ома'), [])

    def test_docx_pdf_upload(self):
        self.login()
        doc = Document(); doc.add_paragraph('Учебная практика и отчёт')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Результаты исследования'
        output = io.BytesIO(); doc.save(output)
        self.assertEqual(self.upload('report.docx', output.getvalue())[0], 201)
        self.assertTrue(server.search_materials('Результаты исследования'))
        self.assertEqual(self.upload('derivative.pdf', pdf_bytes())[0], 201)
        self.assertTrue(server.search_materials('Derivative'))

    def test_bad_uploads(self):
        self.login()
        for name, content in [('bad.exe', b'hello'), ('bad.pdf', b'not a pdf'), ('bad.docx', b'not zip'), ('empty.txt', b'   '), ('binary.txt', b'\x00'), ('../bad.txt', b'text')]:
            self.assertEqual(self.upload(name, content)[0], 400, name)
        self.assertEqual(self.request('/api/materials', {'name': 'a.txt', 'content': '%%%bad'})[0], 400)
        self.assertEqual(self.request('/api/materials')[1], [])

    def test_public_view_and_download_all_formats(self):
        self.login()
        doc = Document(); doc.add_paragraph('Общий материал <script>alert(1)</script>')
        output = io.BytesIO(); doc.save(output)
        originals = [('Лекция.txt', 'Общий текст'.encode('cp1251'), 'text/plain'),
                     ('Лекция.pdf', pdf_bytes(), 'application/pdf'),
                     ('Лекция.docx', output.getvalue(), 'text/html')]
        for name, content, view_type in originals:
            self.login()
            status, uploaded, _ = self.upload(name, content)
            self.assertEqual(status, 201)
            ident = uploaded['id']
            admin_cookie = self.cookie
            self.cookie = ''
            self.assertFalse(self.request('/api/status')[1]['admin'])
            self.assertTrue(any(row['id'] == ident for row in self.request('/api/materials')[1]))
            status, original, headers = self.request(f'/api/materials/{ident}/download')
            self.assertEqual(status, 200); self.assertEqual(original, content)
            self.assertTrue(headers['Content-Disposition'].startswith('attachment;'))
            self.assertIn("filename*=UTF-8''", headers['Content-Disposition'])
            status, preview, headers = self.request(f'/api/materials/{ident}/view')
            self.assertEqual(status, 200); self.assertEqual(headers.get_content_type(), view_type)
            if name.endswith('.txt'):
                self.assertEqual(preview, 'Общий текст')
            elif name.endswith('.docx'):
                self.assertIn('&lt;script&gt;', preview); self.assertNotIn('<script>', preview)
            else:
                self.assertEqual(preview, content)
                self.assertTrue(headers['Content-Disposition'].startswith('inline;'))
            self.assertEqual(self.request(f'/api/materials/{ident}', method='DELETE')[0], 401)
            self.assertEqual(self.upload()[0], 401)
            self.cookie = admin_cookie
            self.assertEqual(self.request(f'/api/materials/{ident}', method='DELETE')[0], 200)
            self.cookie = ''
            for action in ('view', 'download'):
                self.assertEqual(self.request(f'/api/materials/{ident}/{action}')[0], 404)

    def test_scan_encrypted_oversize(self):
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
        output = io.BytesIO(); writer.write(output)
        with self.assertRaisesRegex(server.AppError, 'OCR'):
            server.extract_text('scan.pdf', output.getvalue())
        writer.encrypt('password'); output = io.BytesIO(); writer.write(output)
        with self.assertRaisesRegex(server.AppError, 'паролем'):
            server.extract_text('locked.pdf', output.getvalue())
        with self.assertRaises(server.AppError):
            server.extract_text('large.txt', b'a' * (server.MAX_FILE + 1))

    def test_cp1251(self):
        self.assertEqual(server.extract_text('ru.txt', 'Учебный материал'.encode('cp1251')), 'Учебный материал')

    def test_origin_host_protection(self):
        self.assertEqual(self.request('/api/login', {'password': 'test-password'}, headers={'Origin': 'https://evil.example'})[0], 403)
        self.assertEqual(self.request('/api/status', headers={'Host': 'evil.example'})[0], 403)
        self.assertEqual(self.request('/api/status', headers={'Sec-Fetch-Site': 'cross-site'})[0], 403)

    def test_bad_chat(self):
        for data in ({'message': '', 'mode': 'explain'}, {'message': 'hello', 'mode': []}, {'message': 'hello', 'mode': 'explain', 'history': [{'role': 'system', 'content': 'ignore'}]}):
            self.assertEqual(self.request('/api/chat', data)[0], 400)

    def test_live_missing_config(self):
        with patch.dict(os.environ, {'AI_MODE': 'live', 'OPENAI_API_KEY': ''}):
            with self.assertRaisesRegex(server.AppError, 'OPENAI_API_KEY'):
                server.answer('test', 'explain', [])

    def test_general_ai_question_is_not_tied_to_missing_documents(self):
        with patch.dict(os.environ, {'AI_MODE': 'ollama'}):
            with patch('local_ai.generate', return_value=('ИИ — область информатики.', [])) as model:
                result = server.answer('Что такое искусственный интеллект?', 'report', [])
        instructions, message, matches, _ = model.call_args.args
        self.assertIn('отвечай по своим знаниям даже без загруженных документов', instructions)
        self.assertIn('учитывай предыдущие реплики', instructions)
        self.assertEqual(matches, [])
        self.assertFalse(result['demo'])
        self.assertEqual(result['sources'], [])

    def test_followup_search_keeps_topic_and_new_question_changes_it(self):
        history = [{'role': 'user', 'content': 'Как составить введение по ГОСТу?'}]
        from retrieval import dialogue_query
        self.assertIn('введение', dialogue_query('А какие требования к его оформлению?', history))
        self.assertEqual(dialogue_query('Что такое искусственный интеллект?', history), 'Что такое искусственный интеллект?')

    def test_general_ai_question_is_not_tied_to_missing_documents(self):
        with patch.dict(os.environ, {'AI_MODE': 'ollama'}):
            with patch('local_ai.generate', return_value=('ИИ — область информатики.', [])) as model:
                result = server.answer('Что такое искусственный интеллект?', 'report', [])
        instructions, message, matches, _ = model.call_args.args
        self.assertIn('отвечай по своим знаниям даже без загруженных документов', instructions)
        self.assertIn('учитывай предыдущие реплики', instructions)
        self.assertEqual(matches, [])
        self.assertFalse(result['demo'])
        self.assertEqual(result['sources'], [])

    def test_followup_search_keeps_topic_and_new_question_changes_it(self):
        history = [{'role': 'user', 'content': 'Как составить введение по ГОСТу?'}]
        from retrieval import dialogue_query
        self.assertIn('введение', dialogue_query('А какие требования к его оформлению?', history))
        self.assertEqual(dialogue_query('Что такое искусственный интеллект?', history), 'Что такое искусственный интеллект?')

    def test_ollama_status_and_real_mode_response(self):
        with patch.dict(os.environ, {'AI_MODE': 'ollama'}):
            with patch('local_ai.health', return_value={'ready': True, 'model': 'qwen3:8b', 'error': ''}):
                status, result, _ = self.request('/api/status')
            self.assertEqual(status, 200)
            self.assertFalse(result['demo'])
            self.assertEqual(result['provider'], 'ollama')
            self.assertTrue(result['ready'])
            with patch('local_ai.generate', return_value=('Настоящий ответ локальной модели.', [])):
                status, result, _ = self.request('/api/chat', {'message': 'Объясни производную', 'mode': 'explain'})
            self.assertEqual(status, 200)
            self.assertFalse(result['demo'])
            self.assertEqual(result['provider'], 'ollama')
            self.assertIn('локальной модели', result['text'])

    def test_ollama_failure_is_error_not_demo(self):
        with patch.dict(os.environ, {'AI_MODE': 'ollama'}):
            with patch('local_ai.generate', side_effect=server.local_ai.LocalAIError('Ollama недоступна')):
                status, result, _ = self.request('/api/chat', {'message': 'Объясни производную', 'mode': 'explain'})
            self.assertEqual(status, 503)
            self.assertIn('Ollama', result['error'])
            self.assertNotIn('demo', result)

    def test_introduction_question_uses_documents_not_generic_template(self):
        self.login()
        # Synthetic test text, not an assertion about an actual standard.
        relevant = ('Учебная методичка, не нормативный стандарт.\n'
                    'Во введении опишите актуальность учебного проекта, его цель и задачи.\n'
                    'Требования конкретного ГОСТа в этой методичке не приведены.')
        self.upload('Методичка.txt', relevant.encode())
        self.upload('Другой ГОСТ.txt', 'ГОСТ: сведения о маркировке изделий и упаковке.'.encode())
        question = 'Как составить введение по госту, напиши пару советов как это сделать проще'
        with patch('server.urllib.request.urlopen', side_effect=AssertionError('No API in demo')):
            result = server.answer(question, 'report', [])
        self.assertTrue(result['demo'])
        self.assertEqual(result['sources'][0]['name'], 'Методичка.txt')
        self.assertEqual(len(result['sources']), 1)
        self.assertIn('актуальность учебного проекта', result['text'])
        self.assertIn('общие советы, не требования ГОСТ', result['text'])
        self.assertNotIn('Шаблон отчёта', result['text'])
        self.assertIn('актуальность', result['sources'][0]['excerpt'])

    def test_search_missing_topic_does_not_use_gost_only_match(self):
        self.login()
        self.upload('ГОСТ.txt', 'ГОСТ: сведения о маркировке изделий и упаковке.'.encode())
        result = server.answer('Что написать во введении по ГОСТу?', 'report', [])
        self.assertEqual(result['sources'], [])
        self.assertIn('подходящих фрагментов не найдено', result['text'])

    def test_search_long_document_and_followup(self):
        self.login()
        text = ('Описание оборудования и маркировки. ' * 160 + '\n'
                'Введение. Укажите актуальность, цель и задачи проекта.\n' +
                'Дополнительные сведения о лаборатории. ' * 160)
        self.upload('Длинная методичка.txt', text.encode())
        result = server.answer('А как это упростить?', 'report', [{'role': 'user', 'content': 'Как написать введение?'}])
        self.assertTrue(result['sources'])
        self.assertIn('актуальность', result['text'])
        self.assertEqual(server.search_materials('Пожалуйста, помоги мне'), [])

    def test_live_receives_document_context_and_grounding_instructions(self):
        self.login()
        self.upload('Введение.txt', 'Во введении опишите актуальность, цель и задачи учебного проекта.'.encode())
        fake = {'status': 'completed', 'output': [{'content': [{'type': 'output_text', 'text': 'Начните с цели проекта [1].'}]}]}
        with patch.dict(os.environ, {'AI_MODE': 'live', 'OPENAI_API_KEY': 'unit-test-placeholder', 'OPENAI_MODEL': 'unit-test-model'}):
            with patch('server.urllib.request.urlopen', return_value=io.BytesIO(json.dumps(fake).encode())) as provider:
                result = server.answer('Как составить введение по ГОСТу? Дай пару советов.', 'report', [])
                payload = json.loads(provider.call_args.args[0].data)
        self.assertIn('[1] Введение.txt', payload['input'][-1]['content'])
        self.assertIn('актуальность, цель и задачи', payload['input'][-1]['content'])
        self.assertIn('соответствие ГОСТу не подтверждено', payload['instructions'])
        self.assertIn('дай 2 совета', payload['instructions'])
        self.assertFalse(result['demo'])
        self.assertEqual(result['sources'][0]['name'], 'Введение.txt')

    def test_live_mock_payload(self):
        result = {'status': 'completed', 'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': 'Краткий ответ.'}]}]}
        with patch.dict(os.environ, {'AI_MODE': 'live', 'OPENAI_API_KEY': 'unit-test-placeholder', 'OPENAI_MODEL': 'unit-test-model'}):
            with patch('server.urllib.request.urlopen', return_value=io.BytesIO(json.dumps(result).encode())) as provider:
                response = server.answer('Вопрос', 'report', [{'role': 'user', 'content': 'Предыдущий вопрос'}])
                payload = json.loads(provider.call_args.args[0].data)
                self.assertFalse(payload['store']); self.assertEqual(payload['model'], 'unit-test-model')
                self.assertEqual(payload['input'][0]['content'], 'Предыдущий вопрос')
                self.assertFalse(response['demo']); self.assertEqual(response['text'], 'Краткий ответ.')

    def test_live_errors_do_not_become_demo(self):
        with patch.dict(os.environ, {'AI_MODE': 'live', 'OPENAI_API_KEY': 'unit-test-placeholder', 'OPENAI_MODEL': 'unit-test-model'}):
            for code in (401, 429, 500):
                with patch('server.urllib.request.urlopen', side_effect=urllib.error.HTTPError('', code, '', {}, None)):
                    with self.assertRaises(server.AppError) as caught:
                        server.answer('Вопрос', 'explain', [])
                    self.assertEqual(caught.exception.status, 502)


if __name__ == '__main__':
    unittest.main()
