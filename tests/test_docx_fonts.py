import io
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import docx_review
from docx_fonts import FontResolver


class DocxFontTests(unittest.TestCase):
    def test_cyrillic_uses_hansi_calibri_instead_of_ascii_courier(self):
        doc = Document()
        doc.styles['Normal'].font.name = 'Times New Roman'
        doc.styles['Normal'].font.size = Pt(14)
        p = doc.add_paragraph('Основной текст исследования. ' * 10)
        run = p.add_run('ПИЗДОЙ')
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:ascii'), 'Courier New')
        fonts.set(qn('w:hAnsi'), 'Calibri')
        run.font.size = Pt(20)
        stream = io.BytesIO(); doc.save(stream)
        content = stream.getvalue()
        review, _ = docx_review.analyze('cyrillic.docx', content, 'Проверь')
        comment = review['format_comments'][0]
        self.assertEqual(comment['quote'], 'ПИЗДОЙ')
        self.assertIn('Calibri', comment['text'])
        self.assertNotIn('Courier New', comment['text'])
        output = docx_review.export_comments('cyrillic.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Проверка.', 'demo': True})
        exported = Document(io.BytesIO(output))
        note = next(c.text for c in exported.comments if 'Наблюдение программы' in c.text)
        self.assertIn('Calibri', note)
        self.assertNotIn('Courier New', note)

    def test_same_run_has_distinct_ascii_and_cyrillic_fonts(self):
        doc = Document()
        run = doc.add_paragraph().add_run('ABCАБВ')
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:ascii'), 'Courier New'); fonts.set(qn('w:hAnsi'), 'Calibri')
        spans = FontResolver(doc).spans(run, [])
        self.assertEqual([(text, font) for text, font, _, _ in spans], [('ABC', 'Courier New'), ('АБВ', 'Calibri')])

    def test_theme_and_document_defaults_resolve_calibri(self):
        doc = Document()
        from lxml import etree
        theme = FontResolver(doc).theme
        a = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
        minor = theme.find('.//' + a + 'fontScheme/' + a + 'minorFont')
        minor.find(a + 'latin').set('typeface', 'Calibri')
        for font in minor.findall(a + 'font'):
            if font.get('script') == 'Cyrl':
                font.set('typeface', 'Calibri')
        next(rel.target_part for rel in doc.part.rels.values() if rel.reltype.endswith('/theme'))._blob = etree.tostring(theme)
        run = doc.add_paragraph().add_run('Кириллица')
        spans = FontResolver(doc).spans(run, [doc.styles['Normal']])
        self.assertEqual(spans[0][1:3], ('Calibri', 11.0))

    def test_direct_hansi_overrides_inherited_theme(self):
        doc = Document()
        run = doc.add_paragraph().add_run('Кириллица')
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Arial')
        self.assertEqual(FontResolver(doc).spans(run, [doc.styles['Normal']])[0][1], 'Arial')

    def test_unresolved_theme_does_not_fall_back_to_ascii(self):
        doc = Document()
        run = doc.add_paragraph().add_run('Кириллица')
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:ascii'), 'Courier New'); fonts.set(qn('w:hAnsiTheme'), 'invalid')
        self.assertIsNone(FontResolver(doc).spans(run, [doc.styles['Normal']])[0][1])

    def test_east_asian_hint_for_cyrillic_is_respected(self):
        doc = Document()
        run = doc.add_paragraph().add_run('Кириллица')
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:ascii'), 'Courier New'); fonts.set(qn('w:hAnsi'), 'Arial')
        fonts.set(qn('w:eastAsia'), 'Calibri'); fonts.set(qn('w:hint'), 'eastAsia')
        self.assertEqual(FontResolver(doc).spans(run, [doc.styles['Normal']])[0][1], 'Calibri')

    def test_explicit_complex_script_and_size(self):
        doc = Document()
        run = doc.add_paragraph().add_run('Кириллица')
        properties = run._r.get_or_add_rPr()
        properties.get_or_add_rFonts().set(qn('w:cs'), 'Calibri')
        properties.append(OxmlElement('w:cs'))
        size = OxmlElement('w:szCs'); size.set(qn('w:val'), '30'); properties.append(size)
        self.assertEqual(FontResolver(doc).spans(run, [doc.styles['Normal']])[0][1:3], ('Calibri', 15.0))
