import io
import unittest
import zipfile
import json

from docx import Document
from docx.shared import Cm, Pt

import docx_review


def document_bytes(text='Тестовая работа. Цель — проверить анализ DOCX.'):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    doc.sections[0].left_margin = Cm(3)
    doc.add_heading('Введение', level=1)
    doc.add_paragraph(text)
    doc.add_paragraph('Фрагмент с другим оформлением.').runs[0].font.size = Pt(16)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Табличный факт: ABC-729'
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def context_document_bytes():
    doc = Document()
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(14)
    doc.add_heading('ВВЕДЕНИЕ', level=1)
    p = doc.add_paragraph('Использование технологии WebSocket позволяет перенести управление состоянием игры на сервер, '
                          'обеспечивая мгновенный отклик и честное ')
    r = p.add_run('ПОПОЧКА'); r.font.name = 'Calibri'; r.font.size = Pt(18)
    p.add_run(' баллов. Таким образом, разработка приложения отвечает требованиям производительности, ')
    r = p.add_run('dsjkfgh9uer'); r.font.size = Pt(12)
    p.add_run(' и удобства использования. HTTP и polling описывают способ обмена данными.')
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()


def anchored_text(document, comment):
    from docx.oxml.ns import qn
    active = False
    text = []
    for element in document.part.element.iter():
        if element.tag == qn('w:commentRangeStart') and element.get(qn('w:id')) == str(comment.comment_id):
            active = True
        elif element.tag == qn('w:commentRangeEnd') and element.get(qn('w:id')) == str(comment.comment_id):
            break
        elif active and element.tag == qn('w:t'):
            text.append(element.text or '')
    return ''.join(text)


class DocxReviewTests(unittest.TestCase):
    def test_all_font_outliers_and_random_latin_in_one_paragraph_are_reported(self):
        content = context_document_bytes()
        review, context = docx_review.analyze('context.docx', content, 'Проверь смысл и оформление')
        outliers = review['format_comments']
        self.assertEqual([item['quote'] for item in outliers], ['ПОПОЧКА', 'dsjkfgh9uer'])
        self.assertIn('12 пт', outliers[1]['text'])
        self.assertIn('14 пт', outliers[1]['text'])
        self.assertEqual([item['quote'] for item in review['language_comments']], ['dsjkfgh9uer'])
        self.assertIn('dsjkfgh9uer', context)
        output = docx_review.export_comments('context.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Программная проверка.', 'demo': True})
        checked = Document(io.BytesIO(output))
        notes = [c for c in checked.comments if 'Наблюдение программы' in c.text]
        self.assertEqual([anchored_text(checked, c) for c in notes], ['ПОПОЧКА', 'dsjkfgh9uer'])
        language = next(c for c in checked.comments if 'Проверка языка · абзац' in c.text)
        self.assertEqual(anchored_text(checked, language), 'dsjkfgh9uer')

    def test_language_rules_preserve_terms_urls_code_and_english_prose(self):
        doc = Document()
        doc.add_paragraph('Работа использует WebSocket, polling, HTTP, PostgreSQL, JavaScript и идентификатор dsjkfgh9uer. '
                          'Документация доступна по адресу https://example.org/dsjkfgh9uer и пример `dsjkfgh9uer` является кодом.')
        stream = io.BytesIO(); doc.save(stream)
        review, _ = docx_review.analyze('terms.docx', stream.getvalue(), 'Проверь')
        self.assertEqual(review['language_comments'], [])
        doc = Document(); doc.add_paragraph('This paragraph is entirely in English and contains an identifier dsjkfgh9uer.')
        stream = io.BytesIO(); doc.save(stream)
        review, _ = docx_review.analyze('english.docx', stream.getvalue(), 'Check')
        self.assertEqual(review['language_comments'], [])

    def test_profanity_and_font_outlier_are_detected_and_anchored_to_word(self):
        doc = Document()
        doc.styles['Normal'].font.name = 'Times New Roman'
        doc.styles['Normal'].font.size = Pt(14)
        p = doc.add_paragraph('Основная часть работы и обоснование актуальности. ' * 5)
        changed = p.add_run('ПИЗДОЙ')
        changed.font.name = 'Arial'; changed.font.size = Pt(20)
        p.add_run(' Далее продолжается обычный учебный текст.')
        stream = io.BytesIO(); doc.save(stream)
        content = stream.getvalue()
        review, _ = docx_review.analyze('review.docx', content, 'Проверь')
        self.assertEqual(review['content_comments'][0]['quote'], 'ПИЗДОЙ')
        self.assertEqual(review['content_comments'][0]['category'], 'important')
        self.assertEqual(review['format_comments'][0]['quote'], 'ПИЗДОЙ')
        self.assertIn('Arial', review['format_comments'][0]['text'])
        self.assertIn('20 пт', review['format_comments'][0]['text'])
        self.assertIn('Times New Roman', review['format_comments'][0]['text'])
        output = docx_review.export_comments('review.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Проверка.', 'demo': True})
        checked = Document(io.BytesIO(output))
        self.assertEqual(checked.paragraphs[0].text, p.text)
        comment = next(c for c in checked.comments if 'Проверка лексики' in c.text)
        self.assertEqual(anchored_text(checked, comment), 'ПИЗДОЙ')
        self.assertEqual(comment.paragraphs[0].text, 'Важные замечания и ошибки')

    def test_word_inside_one_run_is_split_without_changing_text_or_style(self):
        doc = Document()
        p = doc.add_paragraph('Начало текста ПИЗДОЙ и продолжение.')
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(14)
        stream = io.BytesIO(); doc.save(stream)
        content = stream.getvalue()
        review, _ = docx_review.analyze('single.docx', content, 'Проверь')
        output = docx_review.export_comments('single.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Проверка.', 'demo': True})
        checked = Document(io.BytesIO(output))
        comment = next(c for c in checked.comments if 'Проверка лексики' in c.text)
        self.assertEqual(anchored_text(checked, comment), 'ПИЗДОЙ')
        self.assertEqual(checked.paragraphs[0].text, p.text)
        for run in checked.paragraphs[0].runs:
            if run.text:
                self.assertEqual(run.font.name, 'Times New Roman')
                self.assertEqual(run.font.size.pt, 14)

    def test_lexical_rules_do_not_match_innocent_words(self):
        doc = Document()
        doc.add_paragraph('Учебный текст. Небо, мебель, художник, хулиган. Подстрахуй коллегу. Исправь плохую формулировку.')
        stream = io.BytesIO(); doc.save(stream)
        review, _ = docx_review.analyze('clean.docx', stream.getvalue(), 'Проверь')
        self.assertEqual(review['content_comments'], [])

    def test_formatting_anchor_uses_actual_run_when_word_is_repeated(self):
        doc = Document()
        doc.styles['Normal'].font.name = 'Times New Roman'
        p = doc.add_paragraph('Текст работы: пример. ' * 10)
        p.add_run('пример').font.name = 'Arial'
        stream = io.BytesIO(); doc.save(stream)
        content = stream.getvalue()
        review, _ = docx_review.analyze('repeated.docx', content, 'Проверь')
        self.assertEqual(review['format_comments'][0]['start'], len(p.text) - len('пример'))
        output = docx_review.export_comments('repeated.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Проверка.', 'demo': True})
        checked = Document(io.BytesIO(output))
        comment = next(c for c in checked.comments if 'Наблюдение программы' in c.text)
        self.assertEqual(anchored_text(checked, comment), 'пример')
    def test_export_uses_three_category_colors_inside_comments_only(self):
        content = document_bytes()
        review, _ = docx_review.analyze('colors.docx', content, 'Проверь')
        comments = [{'paragraph': 2, 'quote': 'Тестовая работа.', 'text': marker, 'category': category}
                    for category, marker in [('recommendation', 'yellow-note'), ('important', 'red-note'), ('improvement', 'green-note')]]
        output = docx_review.export_comments('colors.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Обзор.', 'comments': comments, 'demo': False})
        checked = Document(io.BytesIO(output))
        from docx.oxml.ns import qn
        for category, marker in [('recommendation', 'yellow-note'), ('important', 'red-note'), ('improvement', 'green-note')]:
            comment = next(c for c in checked.comments if marker in c.text)
            style = docx_review.COMMENT_CATEGORIES[category]
            self.assertEqual(comment.paragraphs[0].text, style['label'])
            for paragraph in comment.paragraphs:
                shade = paragraph._p.find(qn('w:pPr')).find(qn('w:shd'))
                self.assertEqual(shade.get(qn('w:fill')), style['fill'])
                for run in paragraph.runs:
                    if run.text:
                        self.assertEqual(str(run.font.color.rgb), style['color'])
        self.assertFalse(checked.part.element.xpath('.//w:shd'))
        original = Document(io.BytesIO(content))
        self.assertEqual([p.text for p in checked.paragraphs], [p.text for p in original.paragraphs])

    def test_old_and_unknown_categories_default_to_recommendations(self):
        document = Document(io.BytesIO(document_bytes()))
        proposed = [{'paragraph': 2, 'quote': 'Тестовая работа.', 'text': 'Проверьте.', 'category': value}
                    for value in [None, 'invalid', ['important']]]
        validated = docx_review.validated_comments(document, proposed)
        self.assertEqual([item['category'] for item in validated], ['recommendation'] * 3)

    def test_export_uses_three_category_colors_inside_comments_only(self):
        content = document_bytes()
        review, _ = docx_review.analyze('colors.docx', content, 'Проверь')
        comments = [{'paragraph': 2, 'quote': 'Тестовая работа.', 'text': marker, 'category': category}
                    for category, marker in [('recommendation', 'yellow-note'), ('important', 'red-note'), ('improvement', 'green-note')]]
        output = docx_review.export_comments('colors.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Обзор.', 'comments': comments, 'demo': False})
        checked = Document(io.BytesIO(output))
        from docx.oxml.ns import qn
        for category, marker in [('recommendation', 'yellow-note'), ('important', 'red-note'), ('improvement', 'green-note')]:
            comment = next(c for c in checked.comments if marker in c.text)
            style = docx_review.COMMENT_CATEGORIES[category]
            self.assertEqual(comment.paragraphs[0].text, style['label'])
            for paragraph in comment.paragraphs:
                shade = paragraph._p.find(qn('w:pPr')).find(qn('w:shd'))
                self.assertEqual(shade.get(qn('w:fill')), style['fill'])
                for run in paragraph.runs:
                    if run.text:
                        self.assertEqual(str(run.font.color.rgb), style['color'])
        self.assertFalse(checked.part.element.xpath('.//w:shd'))
        original = Document(io.BytesIO(content))
        self.assertEqual([p.text for p in checked.paragraphs], [p.text for p in original.paragraphs])

    def test_old_and_unknown_categories_default_to_recommendations(self):
        document = Document(io.BytesIO(document_bytes()))
        proposed = [{'paragraph': 2, 'quote': 'Тестовая работа.', 'text': 'Проверьте.', 'category': value}
                    for value in [None, 'invalid', ['important']]]
        validated = docx_review.validated_comments(document, proposed)
        self.assertEqual([item['category'] for item in validated], ['recommendation'] * 3)

    def test_long_courier_appendix_does_not_flag_times_new_roman_body(self):
        doc = Document()
        doc.styles['Normal'].font.name = 'Times New Roman'
        doc.styles['Normal'].font.size = Pt(14)
        doc.add_heading('Введение', level=1)
        doc.add_paragraph('Основной текст работы: цель, задачи и результаты исследования.')
        doc.add_paragraph('Приложение А')  # Also recognize a manually formatted heading.
        code = doc.add_paragraph('print("appendix")\n' * 500)
        code.runs[0].font.name = 'Courier New'
        code.runs[0].font.size = Pt(10)
        stream = io.BytesIO(); doc.save(stream)
        content = stream.getvalue()
        review, context = docx_review.analyze('appendix.docx', content, 'Проверь оформление')
        self.assertIn('Times New Roman', review['facts'][1])
        self.assertNotIn('Courier New', review['facts'][1])
        self.assertEqual(review['format_comments'], [])
        self.assertTrue(any('Начало приложений' in fact for fact in review['facts']))
        self.assertIn('приложение, отдельное оформление', context)
        output = docx_review.export_comments('appendix.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'Наблюдения по оформлению.', 'demo': True})
        comments = list(Document(io.BytesIO(output)).comments)
        self.assertFalse(any('Наблюдение программы' in comment.text for comment in comments))

    def test_code_styles_are_separate_without_appendix_heading(self):
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        doc.styles['Normal'].font.name = 'Times New Roman'
        doc.styles['Normal'].font.size = Pt(14)
        style = doc.styles.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = 'Courier New'
        style.font.size = Pt(10)
        doc.add_paragraph('Основной текст работы.')
        doc.add_paragraph('print("code")\n' * 300, style='Source Code')
        stream = io.BytesIO(); doc.save(stream)
        review, _ = docx_review.analyze('code.docx', stream.getvalue(), 'Проверь')
        self.assertNotIn('Courier New', review['facts'][1])
        self.assertEqual(review['format_comments'], [])

    def test_contents_and_appendix_mentions_do_not_hide_body_differences(self):
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        doc.styles['Normal'].font.name = 'Times New Roman'
        doc.styles['Normal'].font.size = Pt(14)
        if 'TOC 1' not in doc.styles:
            doc.styles.add_style('TOC 1', WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph('Приложение А .......... 20', style='TOC 1')
        doc.add_paragraph('См. приложение А. Здесь идёт основной текст работы. ' * 10)
        doc.add_paragraph('Ошибка шрифта в основном тексте.').runs[0].font.name = 'Arial'
        stream = io.BytesIO(); doc.save(stream)
        review, _ = docx_review.analyze('contents.docx', stream.getvalue(), 'Проверь')
        self.assertEqual([item['paragraph'] for item in review['format_comments']], [3])
        self.assertFalse(any('Начало приложений' in fact for fact in review['facts']))

    def test_unmarked_courier_in_body_is_not_automatically_exempt(self):
        doc = Document()
        doc.styles['Normal'].font.name = 'Times New Roman'
        doc.add_paragraph('Основная часть: обычный текст исследования. ' * 10)
        doc.add_paragraph('Фрагмент без признаков кода.').runs[0].font.name = 'Courier New'
        stream = io.BytesIO(); doc.save(stream)
        review, _ = docx_review.analyze('body.docx', stream.getvalue(), 'Проверь')
        self.assertEqual([item['paragraph'] for item in review['format_comments']], [2])

    def test_export_preserves_text_formatting_tables_and_existing_comments(self):
        source = Document(io.BytesIO(document_bytes()))
        source.add_comment(source.paragraphs[1].runs, 'Комментарий преподавателя', author='Teacher')
        stream = io.BytesIO(); source.save(stream)
        content = stream.getvalue()
        review, _ = docx_review.analyze('report.docx', content, 'Проверь')
        quote = source.paragraphs[1].text[:20]
        proposed = [{'paragraph': 2, 'quote': quote, 'text': 'Уточните задачи работы.'},
                    {'paragraph': 3, 'quote': 'несуществующая цитата', 'text': 'Общее замечание без привязки.'}]
        result = docx_review.finish_feedback({'text': json.dumps({'summary': 'Краткая проверка.', 'comments': proposed}),
                                              'demo': False, 'document_review': review}, content)
        self.assertEqual(len(review['ai_comments']), 1)
        self.assertIn('Общее замечание без привязки', result['text'])
        self.assertEqual(result['text'].count('Общее замечание без привязки'), 1)
        self.assertEqual(result['text'].count('Общее замечание без привязки'), 1)
        output = docx_review.export_comments('report.docx', content, {
            'source_hash': review['source_hash'], 'summary': result['text'], 'comments': review['ai_comments'], 'demo': False})
        checked = Document(io.BytesIO(output))
        self.assertEqual([p.text for p in checked.paragraphs], [p.text for p in source.paragraphs])
        self.assertEqual(checked.tables[0].cell(0, 0).text, source.tables[0].cell(0, 0).text)
        self.assertEqual(checked.paragraphs[2].runs[0].font.size, source.paragraphs[2].runs[0].font.size)
        self.assertEqual(checked.sections[0].left_margin, source.sections[0].left_margin)
        self.assertTrue(any(c.author == 'Teacher' and c.text == 'Комментарий преподавателя' for c in checked.comments))
        ai = next(c for c in checked.comments if 'Уточните задачи работы' in c.text)
        self.assertEqual(anchored_text(checked, ai), quote)
        from docx.oxml.ns import qn
        ids = [element.get(qn('w:id')) for element in checked.paragraphs[1]._p.findall(qn('w:commentRangeStart'))]
        self.assertIn(str(ai.comment_id), ids)
        self.assertTrue(any('Наблюдение программы' in c.text for c in checked.comments))

    def test_export_rejects_different_source_and_invalid_feedback(self):
        content = document_bytes()
        for feedback in [None, {'source_hash': 'wrong', 'summary': 'Review'}]:
            with self.assertRaises(docx_review.DocumentError):
                docx_review.export_comments('report.docx', content, feedback)

    def test_model_json_sequence_is_recovered_without_guessing_anchors(self):
        content = document_bytes()
        review, _ = docx_review.analyze('report.docx', content, 'Проверь')
        text = json.dumps({'summary': 'Проверка.'}) + ',' + json.dumps({
            'paragraph': 2, 'quote': 'Тестовая работа.', 'text': 'Уточните задачи.'})
        result = docx_review.finish_feedback({'text': text, 'demo': False, 'document_review': review}, content)
        self.assertEqual(result['text'], 'Проверка.')
        self.assertEqual(review['ai_comments'][0]['paragraph'], 2)

    def test_demo_export_does_not_label_client_text_as_ai_review(self):
        content = document_bytes()
        review, _ = docx_review.analyze('report.docx', content, 'Проверь')
        output = docx_review.export_comments('report.docx', content, {
            'source_hash': review['source_hash'], 'summary': 'ДЕМО: параметры файла.', 'demo': True,
            'comments': [{'paragraph': 2, 'quote': 'Тестовая работа.', 'text': 'FAKE-AI-COMMENT'}]})
        comments = '\n'.join(c.text for c in Document(io.BytesIO(output)).comments)
        self.assertIn('ДЕМО', comments)
        self.assertNotIn('FAKE-AI-COMMENT', comments)

    def test_inherited_formatting_tables_and_observations(self):
        review, context = docx_review.analyze('report.docx', document_bytes(), 'Проверь введение')
        facts = '\n'.join(review['facts'])
        for value in ['Times New Roman', '14.0 пт', '16.0 пт', '1.5 строки', '1.25 см', 'слева 3.0']:
            self.assertIn(value, facts)
        self.assertIn('ABC-729', context)
        self.assertIn('Абзац 2', context)
        self.assertTrue(any('может быть намеренное' in note for note in review['notes']))
        self.assertFalse(review['partial'])

    def test_unknown_font_is_not_assumed(self):
        doc = Document(); doc.add_paragraph('Текст с тематическим шрифтом документа.')
        from docx.oxml.ns import qn
        defaults = doc.styles.element.find(qn('w:docDefaults') + '/' + qn('w:rPrDefault') + '/' + qn('w:rPr'))
        defaults.find(qn('w:rFonts')).set(qn('w:hAnsiTheme'), 'unknownTheme')
        stream = io.BytesIO(); doc.save(stream)
        review, _ = docx_review.analyze('theme.docx', stream.getvalue(), 'Проверь')
        self.assertIn('не удалось определить', review['facts'][1])
        self.assertTrue(any('не полностью' in line for line in review['limitations']))

    def test_long_document_has_explicit_partial_scope(self):
        doc = Document()
        for i in range(50):
            doc.add_paragraph(f'Раздел {i}. ' + 'Учебный материал и описание работы. ' * 20)
        doc.add_paragraph('Уникальное заключение ZEBRA729: достигнута цель работы.')
        stream = io.BytesIO(); doc.save(stream)
        review, context = docx_review.analyze('long.docx', stream.getvalue(), 'Проверь заключение ZEBRA729')
        self.assertTrue(review['partial'])
        self.assertIn('ZEBRA729', context)
        self.assertLess(len(context), 8000)

    def test_invalid_files_are_rejected(self):
        for name, content in [('report.doc', b'test'), ('../report.docx', b'test'),
                              ('report.docx', b'not zip'), ('report.docx', b'')]:
            with self.subTest(name=name):
                with self.assertRaises(docx_review.DocumentError):
                    docx_review.analyze(name, content, 'Проверь')

    def test_empty_and_oversized_archives_are_rejected(self):
        stream = io.BytesIO(); Document().save(stream)
        with self.assertRaisesRegex(docx_review.DocumentError, 'не найден текст'):
            docx_review.analyze('empty.docx', stream.getvalue(), 'Проверь')
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('word/document.xml', b'0' * (41 * 1024 * 1024))
        with self.assertRaisesRegex(docx_review.DocumentError, 'распаковки'):
            docx_review.analyze('large.docx', stream.getvalue(), 'Проверь')
