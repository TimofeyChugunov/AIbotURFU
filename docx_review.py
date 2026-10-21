"""Read a private DOCX in memory; report observations, not invented GOST rules."""
import io
import hashlib
import json
import re
import zipfile
from collections import Counter
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from retrieval import terms
from docx_fonts import FontResolver


COMMENT_CATEGORIES = {
    'recommendation': {'label': 'Рекомендации', 'fill': 'FFF2CC', 'color': '54430B'},
    'important': {'label': 'Важные замечания и ошибки', 'fill': 'F4CCCC', 'color': '9C0006'},
    'improvement': {'label': 'Что улучшить', 'fill': 'E2F0D9', 'color': '245529'},
}

# Deliberately narrow whole-word rules; do not match innocent substrings such as
# "учебный", "подстрахуй", "хулиган" or "плохую". This is not a complete dictionary.
PROFANITY = re.compile(
    r'(?:[а-я]*пизд[а-я]*|'
    r'(?:на|по|за|до|ни|не|о|от|с|вы|раз)?ху[йяею][а-я]*|'
    r'(?:за|вы|на|до|от|раз|по|при|под|пере|с|у|об|подъ|съ|въ)?еб(?:а[а-я]*|и[а-я]*|у[а-я]*|ну[а-я]*|е[мшт][а-я]*|л[а-я]*)|'
    r'бля(?:д[а-я]*|ть)?)')


def lexical_comments(paragraphs, scopes):
    comments = []
    for number, (paragraph, _) in enumerate(paragraphs, 1):
        if scopes[number - 1] in ('appendix', 'code', 'contents'):
            continue
        for match in re.finditer(r'[А-Яа-яЁё]+', paragraph.text):
            word = match.group()
            if PROFANITY.fullmatch(word.casefold().replace('ё', 'е')):
                comments.append({'paragraph': number, 'quote': word, 'start': match.start(), 'end': match.end(),
                    'category': 'important', 'text': f'Обнаружена нецензурная лексика: «{word}». '
                    'Для учебного текста замените выражение нейтральным. Если это необходимая цитата, '
                    'проверьте уместность и поясните контекст.'})
                if len(comments) >= 30:
                    return comments
    return comments


TECHNICAL_WORDS = set(('http https websocket websockets polling api json xml html css javascript typescript '
    'python java sql postgresql sqlite mysql docker git github ollama qwen linux windows word docx pdf '
    'utf unicode ascii calibri cambria times new roman courier frontend backend server client request '
    'response token tokens browser chat bot ai rag gpu cpu ram openai fastapi flask django react vue '
    'npm pip localhost tcp udp tls ssl cors csrf oauth jwt nginx redis async await true false null').split())


def language_comments(paragraphs, scopes):
    """Flag likely accidental insertions in Russian prose, not every foreign term."""
    comments = []
    prose = '\n'.join(p.text for (p, _), scope in zip(paragraphs, scopes) if scope in ('body', 'heading', 'table'))
    document_cyrillic = len(re.findall(r'[А-Яа-яЁё]', prose))
    document_latin = len(re.findall(r'[A-Za-z]', prose))
    russian_document = document_cyrillic >= 40 and document_cyrillic >= document_latin
    for number, (paragraph, _) in enumerate(paragraphs, 1):
        if scopes[number - 1] not in ('body', 'heading', 'table'):
            continue
        text = paragraph.text
        cyrillic = len(re.findall(r'[А-Яа-яЁё]', text))
        latin = len(re.findall(r'[A-Za-z]', text))
        if (cyrillic < 20 or cyrillic < latin) and not russian_document:
            continue
        protected = [match.span() for match in re.finditer(r'https?://\S+|\b\S+@\S+\.\S+|`[^`]+`', text)]
        for match in re.finditer(r'(?<![\w/])[A-Za-z][A-Za-z0-9_-]{5,}(?!\w)', text):
            word = match.group()
            if any(start <= match.start() < end for start, end in protected):
                continue
            if word.casefold() in TECHNICAL_WORDS or re.fullmatch(r'[A-Z][A-Z0-9_-]+', word):
                continue
            if re.search(r'[a-z][A-Z]', word) or '_' in word or '-' in word:
                continue  # Likely an identifier or compound term.
            context = text[max(0, match.start() - 65):match.start()].casefold()
            if re.search(r'(?:переменн\w*|идентификатор\w*|команд\w*|хеш\w*|код)\s*[:=«"\s]*$', context):
                continue
            letters = re.sub('[^a-z]', '', word.casefold())
            suspicious = (bool(re.search('[bcdfghjklmnpqrstvwxz]{5,}', letters)) or
                          len(letters) >= 8 and not re.search('[aeiouy]', letters))
            if not suspicious:
                continue
            comments.append({'paragraph': number, 'quote': word, 'start': match.start(), 'end': match.end(),
                'category': 'recommendation', 'text': f'Подозрительная последовательность «{word}» в русскоязычном тексте. '
                'Проверьте случайную вставку, раскладку и смысл предложения. Если это технический термин '
                'или идентификатор, поясните его; сам факт использования латиницы не является ошибкой.'})
            if len(comments) >= 30:
                return comments
    return comments


def anchor_runs(paragraph, item):
    """Select an exact plain-text fragment, retaining run formatting when splitting."""
    runs = [run for run in paragraph.runs if run.text]
    text = ''.join(run.text for run in runs)
    start, end = item.get('start'), item.get('end')
    if (text != paragraph.text or type(start) is not int or type(end) is not int
            or not 0 <= start < end <= len(text) or text[start:end] != item['quote']):
        return [run for run in runs if run.text.strip()]

    def split_at(position):
        offset = 0
        for run in paragraph.runs:
            next_offset = offset + len(run.text)
            if offset < position < next_offset:
                if any(child.tag not in (qn('w:rPr'), qn('w:t')) for child in run._r):
                    return False  # Do not rewrite fields, drawings or other complex runs.
                original = run.text
                right_xml = deepcopy(run._r)
                run.text = original[:position - offset]
                run._r.addnext(right_xml)
                Run(right_xml, paragraph).text = original[position - offset:]
                return True
            offset = next_offset
        return True

    split_at(end)
    split_at(start)
    selected, offset = [], 0
    for run in paragraph.runs:
        next_offset = offset + len(run.text)
        if offset < end and next_offset > start and run.text:
            selected.append(run)
        offset = next_offset
    return selected or runs


def comment_category(value):
    return value if isinstance(value, str) and value in COMMENT_CATEGORIES else 'recommendation'


def add_styled_comment(document, runs, text, category='recommendation'):
    """Shade comment content, without depending on Word's author/balloon colors."""
    style = COMMENT_CATEGORIES[comment_category(category)]
    comment = document.add_comment(runs, '', author='AIbotURFU', initials='AI')
    comment.paragraphs[0].add_run(style['label']).bold = True
    comment.add_paragraph(text)
    for paragraph in comment.paragraphs:
        shade = OxmlElement('w:shd')
        shade.set(qn('w:val'), 'clear')
        shade.set(qn('w:fill'), style['fill'])
        paragraph._p.get_or_add_pPr().append(shade)
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            if not run.text:
                continue
            run.font.color.rgb = RGBColor.from_string(style['color'])
            shade = OxmlElement('w:shd')
            shade.set(qn('w:val'), 'clear')
            shade.set(qn('w:fill'), style['fill'])
            run._r.get_or_add_rPr().append(shade)
    return comment


class DocumentError(ValueError):
    pass


def style_chain(style):
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def inherited(objects, attribute):
    for obj in objects:
        value = getattr(obj, attribute)
        if value is not None:
            return value
    return None


def distribution(counter, unit=''):
    if not counter:
        return 'не удалось определить'
    return ', '.join(f'{value}{unit} ({count})' for value, count in counter.most_common(5))


def collect_paragraphs(document):
    paragraphs, seen_cells = [], set()
    table_count = 0
    nested_skipped = False

    def collect(container, label='', depth=0):
        nonlocal table_count, nested_skipped
        for block in container.iter_inner_content():
            if isinstance(block, Paragraph) and block.text.strip():
                paragraphs.append((block, label))
                if len(paragraphs) > 4000:
                    raise DocumentError('В DOCX больше 4000 непустых абзацев. Разделите работу на части.')
            elif isinstance(block, Table):
                table_count += 1
                if depth >= 3:
                    nested_skipped = True
                    continue
                for row_index, row in enumerate(block.rows, 1):
                    for cell_index, cell in enumerate(row.cells, 1):
                        if cell._tc not in seen_cells:
                            seen_cells.add(cell._tc)
                            collect(cell, f'таблица {table_count}, строка {row_index}, ячейка {cell_index}', depth + 1)
    collect(document)
    return paragraphs, table_count, nested_skipped


def code_style(style):
    name = re.sub(r'[\s_-]+', '', style.name.casefold())
    return name.startswith(('code', 'sourcecode', 'листинг', 'исходныйкод')) or name in {
        'код', 'htmlpreformatted', 'preformattedtext'}


def paragraph_scopes(paragraphs):
    """Appendices are separate from prose; a monospace font alone is not proof of code."""
    scopes = []
    appendix = False
    for paragraph, label in paragraphs:
        styles = list(style_chain(paragraph.style))
        text = paragraph.text.strip()
        heading = any(s.name.lower().startswith(('heading', 'заголовок', 'title')) for s in styles)
        toc = any(s.name.casefold().startswith(('toc', 'содержание', 'оглавление')) for s in styles)
        # Ignore TOC leaders and prose such as "см. приложение А".
        leader = re.search(r'\.{2,}|…|(?:\.\s*){3,}|\t[^\n]*\d+\s*$', text)
        appendix_title = re.fullmatch(
            r'(?:приложение|приложения|appendix|appendices)'
            r'(?:\s+[а-яёa-z0-9])?'
            r'(?:\s*(?:[.\-—–:]\s*|\n|\().*)?', text, re.IGNORECASE)
        styled_appendix = heading and re.match(r'^(?:приложени[ея]|appendix)\s+[а-яёa-z0-9]\b', text, re.IGNORECASE)
        if not label and not toc and not leader and len(text) <= 180 and (appendix_title or styled_appendix):
            appendix = True
        runs = [run for run in paragraph.runs if run.text.strip()]
        all_code = bool(runs) and all(any(code_style(s) for s in style_chain(run.style)) for run in runs)
        scope = ('contents' if toc else 'appendix' if appendix else 'heading' if heading else
                 'table' if label else 'code' if all_code or any(code_style(s) for s in styles) else 'body')
        scopes.append(scope)
    return scopes


def analyze(name, content, question):
    if not isinstance(name, str) or not name.lower().endswith('.docx'):
        raise DocumentError('В чат можно прикрепить только DOCX. Сохраните DOC или PDF как DOCX.')
    if not 1 <= len(name) <= 180 or any(c in name for c in '/\\') or any(ord(c) < 32 for c in name):
        raise DocumentError('Некорректное имя файла.')
    if not content or len(content) > 10 * 1024 * 1024:
        raise DocumentError('Размер DOCX должен быть от 1 байта до 10 МБ.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 40 * 1024 * 1024:
                raise DocumentError('DOCX слишком большой после распаковки: предел 40 МБ и 2000 элементов.')
            if any(item.filename.lower().endswith('vbaproject.bin') for item in entries):
                raise DocumentError('Документы с макросами не поддерживаются. Сохраните обычный DOCX.')
        document = Document(io.BytesIO(content))
    except DocumentError:
        raise
    except Exception:
        raise DocumentError('Не удалось прочитать DOCX. Проверьте формат и целостность файла.')

    paragraphs, table_count, nested_skipped = collect_paragraphs(document)
    scopes = paragraph_scopes(paragraphs)
    scope_names = {'body': 'основной текст', 'appendix': 'приложение, отдельное оформление', 'code': 'листинг кода, отдельное оформление',
                   'heading': 'заголовок', 'contents': 'оглавление', 'table': 'таблица'}
    texts = [f'Абзац {i} ({", ".join(filter(None, [scope_names.get(scopes[i - 1]), label]))}): {p.text.strip()}'
             for i, (p, label) in enumerate(paragraphs, 1)]
    characters = sum(len(p.text) for p, _ in paragraphs)
    if not characters:
        raise DocumentError('В DOCX не найден текст. Изображения и сканы пока не распознаются.')
    if characters > 500_000:
        raise DocumentError('В документе больше 500 000 символов. Разделите работу на части.')

    fonts, sizes, spacing, indents = Counter(), Counter(), Counter(), Counter()
    body_signatures = []
    unknown_runs = 0
    headings = []
    font_resolver = FontResolver(document)
    for i, (paragraph, label) in enumerate(paragraphs, 1):
        styles = list(style_chain(paragraph.style))
        heading = any(style.name.lower().startswith(('heading', 'заголовок', 'title')) for style in styles)
        if heading:
            headings.append(paragraph.text[:120])
        body = scopes[i - 1] == 'body'
        if body:
            formats = [paragraph.paragraph_format, *(s.paragraph_format for s in styles)]
            line = inherited(formats, 'line_spacing')
            if line is not None:
                spacing[f'{round(line.pt, 2)} пт' if hasattr(line, 'pt') else f'{line:g} строки'] += 1
            indent = inherited(formats, 'first_line_indent')
            if indent is not None:
                indents[round(indent.cm, 2)] += 1
        run_offset = 0
        for run in paragraph.runs:
            run_start = run_offset
            run_offset += len(run.text)
            if not run.text.strip():
                continue
            run_styles = list(style_chain(run.style))
            if any(code_style(style) for style in run_styles):
                continue
            if body:
                for text, font, size, offset in font_resolver.spans(run, [*run_styles, *styles]):
                    if not any(character.isalnum() for character in text):
                        continue
                    if font is not None:
                        fonts[font] += len(text)
                    if size is not None:
                        sizes[size] += len(text)
                    if font is None or size is None:
                        unknown_runs += 1
                    body_signatures.append((i, font, size, text, run_start + offset))

    facts = [f'Непустых абзацев: {len(paragraphs)}; таблиц: {table_count}; символов: {characters}.',
             'Шрифты основного текста (число символов): ' + distribution(fonts) + '.',
             'Размеры основного текста (число символов): ' + distribution(sizes, ' пт') + '.',
             'Межстрочный интервал (число абзацев): ' + distribution(spacing) + '.',
             'Отступ первой строки (число абзацев): ' + distribution(indents, ' см') + '.']
    excluded = Counter(scope for scope in scopes if scope != 'body')
    if excluded:
        facts.append('Из сравнения основного оформления исключены: ' + ', '.join(
            f'{scope_names[scope]} — {count} абз.' for scope, count in excluded.items()) + '.')
    first_appendix = next((i for i, scope in enumerate(scopes, 1) if scope == 'appendix'), None)
    if first_appendix:
        facts.append(f'Начало приложений распознано в абзаце {first_appendix}. '
                     'Их шрифты, размеры, интервалы и отступы не задают образец для основной части.')
    facts.append('Преобладающее оформление определяется только по основной части и не является нормативом. '
                 'Другой шрифт в коде или приложении сам по себе не означает ошибку.')
    for i, section in enumerate(document.sections, 1):
        if i > 10:
            facts.append('Показаны поля первых 10 разделов документа.')
            break
        def cm(value):
            return str(round(value.cm, 2)) if value is not None else 'не определено'
        facts.append(f'Раздел {i}: поля слева {cm(section.left_margin)}, справа {cm(section.right_margin)}, '
                     f'сверху {cm(section.top_margin)}, снизу {cm(section.bottom_margin)} см.')
    notes = []
    format_comments = []
    if len(fonts) > 1 or len(sizes) > 1:
        main_font = fonts.most_common(1)[0][0] if fonts else None
        main_size = sizes.most_common(1)[0][0] if sizes else None
        different = sorted({i for i, font, size, _, _ in body_signatures
                            if (font is not None and main_font and font != main_font)
                            or (size is not None and main_size and size != main_size)})
        if different:
            notes.append('Различия шрифта или размера в основном тексте: абзацы '
                         + ', '.join(map(str, different[:12]))
                         + '. Проверьте единообразие; это может быть намеренное оформление, а не ошибка.')
            seen_fragments = set()
            for number, font, size, fragment, offset in body_signatures:
                if not ((font is not None and main_font and font != main_font)
                        or (size is not None and main_size and size != main_size)):
                    continue
                identity = (number, offset, font, size)
                if identity in seen_fragments:
                    continue
                seen_fragments.add(identity)
                start = offset + len(fragment) - len(fragment.lstrip())
                quote = fragment.strip()[:120]
                details = []
                if font is not None and main_font and font != main_font:
                    details.append(f'шрифт {font}, в основной части преобладает {main_font}')
                if size is not None and main_size and size != main_size:
                    details.append(f'размер {size:g} пт, в основной части преобладает {main_size:g} пт')
                format_comments.append({'paragraph': number, 'quote': quote, 'start': start, 'end': start + len(quote),
                    'text': f'Фрагмент «{quote}»: ' + '; '.join(details) + '. '
                            'Проверьте единообразие по методичке. Различие может быть намеренным; это не установленное нарушение ГОСТ.'})
                if len(format_comments) >= 30:
                    break
            for item in format_comments[:5]:
                notes.append(f'Абзац {item["paragraph"]}: {item["text"]}')
    if len(spacing) > 1:
        notes.append('В основном тексте различается межстрочный интервал. Проверьте, оправданы ли различия.')
    if len(indents) > 1:
        notes.append('В основном тексте различается отступ первой строки. Сопоставьте с методичкой.')
    if not headings:
        notes.append('Заголовки стилями Word не обнаружены. Визуальные заголовки могут быть оформлены вручную.')
    content_comments = lexical_comments(paragraphs, scopes)
    for item in content_comments[:5]:
        notes.append(f'Важное замечание · абзац {item["paragraph"]}: {item["text"]}')
    language_findings = language_comments(paragraphs, scopes)
    for item in language_findings[:5]:
        notes.append(f'Проверка языка · абзац {item["paragraph"]}: {item["text"]}')
    body_text = '\n'.join(p.text for (p, _), scope in zip(paragraphs, scopes) if scope == 'body')
    facts.append(f'Состав основной части: кириллических букв {len(re.findall("[А-Яа-яЁё]", body_text))}, '
                 f'латинских букв {len(re.findall("[A-Za-z]", body_text))}. Латиница сама по себе допустима в терминах и идентификаторах.')
    limitations = [
        'Это параметры и наблюдения, а не подтверждение соответствия ГОСТ или методичке.',
        'Не проверены номера страниц, переносы, расположение объектов, изображения, формулы, '
        'колонтитулы, сноски, текстовые поля, исправления и итоговый вид страниц в Word.',
        'Проверяются шрифты, заданные в DOCX с учётом кириллицы, стилей и темы. '
        'Подмена отсутствующего шрифта системой и итоговая отрисовка Word не воспроизводятся. '
        'Часть настроек таблиц, нумерации и сложных систем письма может быть не определена.',
        'Приложения распознаются по отдельным заголовкам, код — по стилям Word. '
        'Если заголовки или стили не обозначены явно, разделение может быть неполным.',
        'Проверка нецензурной лексики использует ограниченные словарные правила для основной части; '
        'замаскированные, редкие и иноязычные выражения могут быть пропущены. Показано до 30 находок.',
        'Показываются до 30 отличающихся фрагментов оформления и до 30 подозрительных языковых вставок. '
        'Проверка языка эвристическая, не заменяет словарь и контекстную редактуру. Логические замечания ИИ требуют проверки человеком.',
    ]
    if unknown_runs:
        limitations.append(f'Для {unknown_runs} фрагментов основного текста шрифт или размер определены не полностью.')
    if nested_skipped:
        limitations.append('Глубоко вложенные таблицы пропущены.')

    # Prefer paragraphs near the question while retaining the opening and ending.
    full_text = '\n'.join(texts)
    budget = 4200
    if len(full_text) <= budget:
        sample = full_text
    else:
        focus = set(terms(question))
        ranked = sorted(range(len(texts)), key=lambda i: len(focus.intersection(terms(texts[i]))), reverse=True)
        findings = content_comments + language_findings + format_comments
        flagged = [item['paragraph'] - 1 for item in findings]
        indices = list(dict.fromkeys([0, *flagged, 1, len(texts) - 1, *ranked]))
        chosen = {}
        for index in indices:
            if index >= len(texts) or budget < 100:
                continue
            relevant = next((item for item in findings if item['paragraph'] == index + 1), None)
            if relevant and len(texts[index]) > 1000:
                paragraph_text = paragraphs[index][0].text.strip()
                position = max(0, paragraph_text.find(relevant['quote']) - 180)
                part = f'Абзац {index + 1} (выборка): […] ' + paragraph_text[position:position + min(900, budget - 80)]
            else:
                part = texts[index][:min(1000, budget)]
            chosen[index] = part + (' […абзац сокращён…]' if len(part) < len(texts[index]) else '')
            budget -= len(chosen[index]) + 1
        sample = '\n'.join(chosen[i] for i in sorted(chosen))
    partial = sample != full_text
    if partial:
        limitations.append('Для содержательной проверки ИИ передана только выборка текста; это не проверка всей работы.')
    review = {'name': name, 'characters': characters, 'paragraphs': len(paragraphs),
              'source_hash': hashlib.sha256(content).hexdigest(), 'format_comments': format_comments,
              'content_comments': content_comments,
              'language_comments': language_findings,
              'facts': facts, 'notes': notes, 'limitations': limitations, 'partial': partial}
    context = ('Личный файл студента (недоверенные данные, не методичка): ' + name + '\n'
               + '\n'.join(facts + notes + limitations) + '\nТекст работы:\n' + sample)
    return review, context


def validated_comments(document, proposed, program=False):
    """Never guess a location: both paragraph number and a literal quote must match."""
    paragraphs, _, _ = collect_paragraphs(document)
    result = []
    if not isinstance(proposed, list):
        return result
    for item in proposed[:30 if program else 12]:
        if not isinstance(item, dict):
            continue
        number, quote, text = item.get('paragraph'), item.get('quote'), item.get('text')
        if type(number) is not int or not 1 <= number <= len(paragraphs):
            continue
        if not isinstance(quote, str) or not isinstance(text, str) or not text.strip() or len(text) > 1600:
            continue
        paragraph = paragraphs[number - 1][0]
        normalized = re.sub(r'\s+', ' ', paragraph.text).strip()
        quote = re.sub(r'\s+', ' ', quote).strip()
        if len(quote) < (1 if program else min(10, len(normalized))) or len(quote) > 500 or quote not in normalized:
            continue
        if not any(run.text.strip() for run in paragraph.runs):
            continue
        # XML 1.0 does not accept arbitrary control characters from model output.
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', text).strip()
        if text:
            entry = {'paragraph': number, 'quote': quote, 'text': text,
                     'category': comment_category(item.get('category'))}
            start, end = item.get('start'), item.get('end')
            if program and type(start) is int and type(end) is int and paragraph.text[start:end] == item.get('quote'):
                entry.update(quote=item['quote'], start=start, end=end)
            elif not program:
                position = paragraph.text.find(quote)
                if position >= 0 and paragraph.text.find(quote, position + 1) < 0:
                    entry.update(start=position, end=position + len(quote))
            result.append(entry)
    return result


def finish_feedback(result, content):
    result = finish_model_feedback(result, content)
    comments = result['document_review'].get('content_comments', [])
    if comments:
        lines = [f'Абзац {item["paragraph"]}: «{item["quote"]}» — нецензурная лексика.' for item in comments[:5]]
        result['text'] = ('Программная проверка: важные замечания\n' + '\n'.join(lines) + '\n\n' + result['text'])[:11500]
    language = result['document_review'].get('language_comments', [])
    if language:
        lines = [f'Абзац {item["paragraph"]}: проверьте смысл вставки «{item["quote"]}».' for item in language[:5]]
        result['text'] = ('Проверка языка: подозрительные вставки\n' + '\n'.join(lines) + '\n\n' + result['text'])[:11500]
    return result


def finish_model_feedback(result, content):
    review = result['document_review']
    review['ai_comments'] = []
    if result.get('demo'):
        return result
    text = result['text'].strip()
    if text.startswith('```'):
        text = re.sub(r'^```(?:json)?\s*|\s*```$', '', text)
    try:
        parsed = json.loads(text)
    except (ValueError, TypeError):
        # Small local models sometimes return a summary object followed by comment
        # objects instead of wrapping the latter in an array. Read real JSON
        # objects; never extract paragraph numbers from arbitrary prose.
        decoder = json.JSONDecoder()
        try:
            parsed, end = decoder.raw_decode(text)
        except ValueError:
            return result  # Plain answers become a clearly labelled general comment.
        if not isinstance(parsed, dict) or not isinstance(parsed.get('summary'), str):
            return result
        extra = []
        tail = text[end:].strip(' ,\r\n\t')
        while tail and len(extra) < 12:
            try:
                item, end = decoder.raw_decode(tail)
            except ValueError:
                review['limitations'].append('Часть ответа ИИ не удалось разобрать как точные замечания.')
                break
            if isinstance(item, dict):
                extra.append(item)
            tail = tail[end:].strip(' ,\r\n\t')
        original = parsed.get('comments', [])
        parsed['comments'] = (original if isinstance(original, list) else []) + extra
    if not isinstance(parsed, dict) or not isinstance(parsed.get('summary'), str) or not parsed['summary'].strip():
        return result
    result['text'] = parsed['summary'].strip()[:6000]
    proposed = parsed.get('comments', [])
    review['ai_comments'] = validated_comments(Document(io.BytesIO(content)), proposed)
    if isinstance(proposed, list) and len(review['ai_comments']) < len(proposed):
        review['limitations'].append('Часть замечаний ИИ не привязана к абзацам: цитаты не совпали с исходным текстом.')
        anchored = {item['text'] for item in review['ai_comments']}
        general = [item['text'][:1600] for item in proposed[:4]
                   if isinstance(item, dict) and isinstance(item.get('text'), str) and item['text'] not in anchored]
        if general:
            result['text'] += '\n\nОбщие замечания без точной привязки:\n' + '\n'.join(general)
            result['text'] = result['text'][:11000]
    return result


def export_comments(name, content, feedback):
    review, _ = analyze(name, content, '')  # Revalidate the file and recompute program observations.
    if not isinstance(feedback, dict) or feedback.get('source_hash') != review['source_hash']:
        raise DocumentError('Это другой файл. Прикрепите исходный DOCX, для которого получен ответ.')
    summary = feedback.get('summary')
    if not isinstance(summary, str) or not summary.strip() or len(summary) > 12000:
        raise DocumentError('Некорректный текст проверки.')
    document = Document(io.BytesIO(content))
    paragraphs, _, _ = collect_paragraphs(document)
    first = next((p for p, _ in paragraphs if any(r.text.strip() for r in p.runs)), None)
    if first is None:
        raise DocumentError('Не удалось привязать комментарии к тексту этого DOCX.')
    general = ('Обзор проверки всего документа (не замечание только к этому абзацу).\n'
               + ('ДЕМО: только программные наблюдения, содержание ИИ не проверялось.\n' if feedback.get('demo') else 'Рекомендации ИИ требуют проверки человеком.\n')
               + summary + '\n\nОграничения:\n' + '\n'.join(review['limitations']))
    general = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', general)
    add_styled_comment(document, [r for r in first.runs if r.text.strip()], general)
    entries = [('Наблюдение программы', item) for item in validated_comments(document, review['format_comments'], program=True)]
    entries += [('Проверка лексики', item) for item in validated_comments(document, review['content_comments'], program=True)]
    entries += [('Проверка языка', item) for item in validated_comments(document, review['language_comments'], program=True)]
    if not feedback.get('demo'):
        entries += [('Рекомендация ИИ', item) for item in validated_comments(document, feedback.get('comments', []))]
    for label, item in entries:
        paragraph = paragraphs[item['paragraph'] - 1][0]
        runs = anchor_runs(paragraph, item)
        add_styled_comment(document, runs, f'{label} · абзац {item["paragraph"]}\n{item["text"]}', item['category'])
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
